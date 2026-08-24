#!/usr/bin/env python3
"""
Génère src/data/romeData.ts à partir des fichiers XLSX officiels ROME (raw/).

Sources :
- 250528-fiches-rome-26m06-tag-pour-diffusion.xlsx  -> 1912 fiches ROME + tags transitions
- 260611-arborescence-simplifiee-des-competences.xlsx -> compétences -> fiches ROME mobilisantes
- 260609-ref-formacode-v14-rome-26m06-v61-open-data.xlsx -> mapping ROME -> FORMACODE (formations suggérées)

Sortie : src/data/romeData.ts (type + 4 exports).
"""
import openpyxl
import warnings
import re
import unicodedata
import os
import json

warnings.filterwarnings("ignore")

RAW = os.path.join(os.path.dirname(__file__), "..", "raw")
OUT = os.path.join(os.path.dirname(__file__), "..", "src", "data", "romeData.ts")


def norm(s: str) -> str:
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = re.sub(r"[^a-z0-9 ]", " ", s.lower())
    s = re.sub(r"\s+", " ", s).strip()
    return s


def main():
    # ---------- 1. Fiches ROME taguées ----------
    wb = openpyxl.load_workbook(os.path.join(RAW, "250528-fiches-rome-26m06-tag-pour-diffusion.xlsx"), read_only=True)
    ws = wb["Tag des fiches ROME"]
    fiches = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row[2] is None:
            continue
        fiches.append({
            "code": str(row[2]).strip(),
            "libelle": str(row[3] or "").strip(),
            "grandDomaine": str(row[0] or "").strip(),
            "domaine": str(row[1] or "").strip(),
            "transitionEcologique": str(row[5] or "").strip(),
            "transitionNumerique": str(row[6] or "").strip(),
            "transitionDemographique": str(row[7] or "").strip(),
            "emploiCadre": str(row[8] or "").strip(),
            "emploiReglemente": str(row[9] or "").strip(),
        })
    wb.close()
    print(f"Fiches ROME taguées : {len(fiches)}")

    # ---------- 2. Arborescence simplifiée des compétences ----------
    wb = openpyxl.load_workbook(os.path.join(RAW, "260611-arborescence-simplifiee-des-competences.xlsx"), read_only=True)
    ws = wb["REF - Comp"]
    # Colonne 'Compétence' = index 13, "Liste des fiches ROME mobilisant l'objet" = index 16
    skill_to_rome: dict[str, set] = {}
    rome_to_skills: dict[str, set] = {}
    n_rows = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        n_rows += 1
        skill = row[13]
        rome_list = row[16]
        if not skill:
            continue
        codes = []
        if rome_list:
            codes = [c.strip() for c in re.split(r"[,\s]+", str(rome_list)) if re.match(r"^[A-Z][0-9]{4}$", c.strip())]
        key = norm(str(skill))
        if not key:
            continue
        if codes:
            skill_to_rome.setdefault(key, set()).update(codes)
            for c in codes:
                rome_to_skills.setdefault(c, set()).add(str(skill).strip())
    wb.close()
    print(f"Lignes compétences lues : {n_rows}")
    print(f"Compétences uniques avec mapping ROME : {len(skill_to_rome)}")
    print(f"Codes ROME présents dans l'arborescence : {len(rome_to_skills)}")

    # Couverture : combien de fiches taguées ont des compétences dans l'arborescence ?
    fiche_codes = {f["code"] for f in fiches}
    covered = fiche_codes & set(rome_to_skills.keys())
    print(f"Fiches taguées couvertes par l'arborescence : {len(covered)}/{len(fiche_codes)}")

    # ---------- 3. Mapping ROME -> FORMACODE ----------
    wb = openpyxl.load_workbook(os.path.join(RAW, "260609-ref-formacode-v14-rome-26m06-v61-open-data.xlsx"), read_only=True)
    ws = wb["ROME 26M06 - Formacode V14"]
    rome_formacode: dict[str, list] = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if row[0] is None or row[4] is None:
            continue
        code = str(row[0]).strip()
        fc_code = str(row[3] or "").strip()
        fc_lib = str(row[4]).strip()
        if not fc_lib:
            continue
        rome_formacode.setdefault(code, []).append({"code": fc_code, "libelle": fc_lib})
    wb.close()
    # Ne garder que 4 formations max par fiche (dédup)
    for c in rome_formacode:
        seen = set()
        dedup = []
        for f in rome_formacode[c]:
            if f["libelle"] in seen:
                continue
            seen.add(f["libelle"])
            dedup.append(f)
        rome_formacode[c] = dedup[:4]
    print(f"Codes ROME avec mapping FORMACODE : {len(rome_formacode)}")

    # ---------- 4. Écriture du module TS ----------
    # Limiter les listes de compétences par fiche (garde ~40 max pour la lisibilité / taille)
    rome_to_skills_trimmed = {c: sorted(list(s))[:40] for c, s in sorted(rome_to_skills.items())}
    skill_to_rome_sorted = {k: sorted(list(v)) for k, v in sorted(skill_to_rome.items())}

    def ts_str_array(items) -> str:
        return "[\n" + ",\n".join("    " + json.dumps(i, ensure_ascii=False) for i in items) + "\n  ]"

    def ts_record_strings(rec: dict) -> str:
        parts = []
        for k, v in sorted(rec.items()):
            parts.append("  " + json.dumps(k, ensure_ascii=False) + ": " + json.dumps(v, ensure_ascii=False))
        return "{\n" + ",\n".join(parts) + "\n}"

    header = """// ============================================================================
// DONNÉES RÉFÉRENTIEL ROME — GÉNÉRÉES AUTOMATIQUEMENT (ne pas éditer à la main)
// Source : scripts/build_rome_data.py + fichiers XLSX officiels France Travail
// (raw/250528-fiches-rome-26m06-tag-pour-diffusion.xlsx,
//  raw/260611-arborescence-simplifiee-des-competences.xlsx,
//  raw/260609-ref-formacode-v14-rome-26m06-v61-open-data.xlsx)
// Généré le : 2026-08-24
// ============================================================================

export interface RomeFiche {
  code: string;
  libelle: string;
  grandDomaine: string;
  domaine: string;
  transitionEcologique: string;
  transitionNumerique: string;
  transitionDemographique: string;
  emploiCadre: string;
  emploiReglemente: string;
}

export const ROME_FICHES: RomeFiche[] = """

    with open(OUT, "w", encoding="utf-8") as f:
        f.write(header)
        f.write(ts_str_array(fiches))
        f.write(";\n\n")

        f.write("// Code ROME -> compétences mobilisées (arborescence simplifiée, max 40 par fiche)\n")
        f.write("export const ROME_CODE_SKILLS: Record<string, string[]> = ")
        f.write(ts_record_strings(rome_to_skills_trimmed))
        f.write(";\n\n")

        f.write("// Compétence normalisée -> codes ROME qui la mobilisent\n")
        f.write("export const SKILL_TO_ROME: Record<string, string[]> = ")
        f.write(ts_record_strings(skill_to_rome_sorted))
        f.write(";\n\n")

        f.write("// Code ROME -> formations FORMACODE suggérées\n")
        f.write("export const ROME_FORMACODE: Record<string, { code: string; libelle: string }[]> = ")
        f.write(ts_record_strings(rome_formacode))
        f.write(";\n")

    size = os.path.getsize(OUT)
    print(f"\nÉcrit : {OUT} ({size/1024:.0f} Ko)")


if __name__ == "__main__":
    main()
